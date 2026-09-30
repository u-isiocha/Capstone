"""Append complete query rounds and charge the associated budget on append."""
from copy import copy
from datetime import datetime, timezone
import json
from pathlib import Path
import shutil
import tempfile

import numpy as np
from openpyxl import load_workbook


def read_budget(workbook):
    path = Path(workbook).with_name("evaluation_budget.json")
    budget = json.loads(path.read_text(encoding="utf-8"))
    used = budget["used_rounds"]
    if (not isinstance(used, int) or not 0 <= used <= budget["total_rounds"]
            or budget["remaining_rounds"] != budget["total_rounds"] - used
            or [r["round"] for r in budget["round_log"]] != list(range(1, used + 1))):
        raise ValueError("Budget counter and round log must be reconciled before querying.")
    return path, budget


def refresh_repository(workbook):
    """Refresh the local repository data snapshot for workspace query requests."""
    workbook = Path(workbook).resolve()
    destination = workbook.parent / "repository" / "data"
    if destination.is_dir():
        for source in (workbook, workbook.with_name("evaluation_budget.json")):
            target = destination / source.name
            if not target.exists() or source.read_bytes() != target.read_bytes():
                shutil.copyfile(source, target)


def append_round(workbook, recommendations):
    """Validate all eight sheets, then save one round and charge it exactly once.

    Stage both files before replacement and roll back on ordinary write errors.
    This is not a crash-proof, multi-process database transaction.
    """
    workbook = Path(workbook).resolve()
    budget_path, budget = read_budget(workbook)
    names = [f"function_{i}" for i in range(1, 9)]
    if set(recommendations) != set(names):
        raise ValueError("A complete eight-function round is required.")
    if budget["remaining_rounds"] == 0:
        return {name: "withheld: evaluation budget exhausted" for name in names}
    original_book = workbook.read_bytes()
    original_budget = budget_path.read_bytes()
    book = load_workbook(workbook)
    try:
        rows = {}
        for name in names:
            sheet = book[name]
            values = np.asarray(recommendations[name].point, dtype=float)
            if (values.shape != (sheet.max_column - 1,)
                    or not np.all(np.isfinite(values))
                    or np.any((values < 0) | (values > 1))):
                raise ValueError(f"{name}: invalid candidate vector")
            if any(row[-1] is None for row in sheet.iter_rows(min_row=2, values_only=True)):
                return {n: f"withheld: {name} has an unfinished row; whole round withheld" for n in names}
            if any(np.allclose(values, row[:-1], rtol=0, atol=1e-9)
                   for row in sheet.iter_rows(min_row=2, values_only=True)):
                return {n: f"withheld: {name} candidate duplicates an existing row; whole round withheld" for n in names}
            # Use explicit round-to-row evidence, not worksheet size to infer usage.
            if budget["round_log"]:
                last_row = budget["round_log"][-1]["workbook_rows"][name]
                if sheet.max_row != last_row:
                    raise ValueError(f"{name}: workbook and budget row mapping disagree")
            rows[name] = sheet.max_row + 1
        for name in names:
            sheet = book[name]
            target_row = rows[name]
            for column in range(1, sheet.max_column + 1):
                source = sheet.cell(target_row - 1, column)
                target = sheet.cell(target_row, column)
                target._style = copy(source._style)
            for column, value in enumerate(recommendations[name].point, 1):
                sheet.cell(target_row, column, float(value))
        number = budget["used_rounds"] + 1
        budget["used_rounds"] = number
        budget["remaining_rounds"] = budget["total_rounds"] - number
        budget["round_log"].append({
            "round": number, "status": "appended_outputs_pending",
            "appended_at": datetime.now(timezone.utc).isoformat(),
            "workbook_rows": rows,
        })
        with tempfile.TemporaryDirectory(dir=workbook.parent, prefix=".round-stage-") as folder:
            staged_book = Path(folder) / workbook.name
            staged_budget = Path(folder) / budget_path.name
            book.save(staged_book)
            staged_budget.write_text(json.dumps(budget, indent=2) + "\n", encoding="utf-8")
            if workbook.read_bytes() != original_book or budget_path.read_bytes() != original_budget:
                raise RuntimeError("Workbook or budget changed during append; no changes saved.")
            staged_book.replace(workbook)
            try:
                staged_budget.replace(budget_path)
            except OSError:
                workbook.write_bytes(original_book)
                raise
        refresh_repository(workbook)
        return {name: f"appended to row {rows[name]}; round {number} charged" for name in names}
    finally:
        book.close()
