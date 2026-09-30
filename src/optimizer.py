"""Recommend the next evaluation for each capstone black-box function.

Uses scikit-learn Gaussian processes with ARD RBF kernels, learned noise,
Sobol global coverage, and trend-aware local refinement.

Preview: python src/optimizer.py --no-append --seed 20260825
"""

from __future__ import annotations

import argparse
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import warnings

import numpy as np
import pandas as pd
from scipy.stats import norm, qmc
from sklearn.exceptions import ConvergenceWarning
from sklearn.gaussian_process import GaussianProcessRegressor
from sklearn.gaussian_process.kernels import ConstantKernel, RBF, WhiteKernel

DEFAULT_SEED = int(datetime.now().timestamp())
DEFAULT_GLOBAL_CANDIDATES = 65_536
DEFAULT_LOCAL_CANDIDATES = 12_000
LOCAL_RADII = (0.20, 0.10, 0.05, 0.025, 0.01)

# Output-only values are benchmarks, never surrogate training rows.
EXTERNAL_OUTPUT_BENCHMARKS: dict[str, float | None] = {
    "function_1": None,
    "function_2": None,
    "function_3": -0.00551,
    "function_4": None,
    "function_5": 2883.2,
    "function_6": None,
    "function_7": None,
    "function_8": 9.9369,
}


@dataclass(frozen=True)
class Policy:
    acquisition: str
    exploration: float
    label: str
    rationale: str


@dataclass(frozen=True)
class Recommendation:
    point: np.ndarray
    predicted_mean: float
    predicted_std: float
    policy: Policy
    trend: str
    kernel: str
    caveat: str


def transform_output(y: np.ndarray, function_number: int) -> np.ndarray:
    if function_number == 5:
        if np.any(y <= 0):
            raise ValueError("Function 5's log transform requires positive outputs.")
        return np.log(y)
    return y.copy()


def inverse_prediction(
    mean: np.ndarray, std: np.ndarray, function_number: int
) -> tuple[np.ndarray, np.ndarray]:
    if function_number != 5:
        return mean, std
    original_mean = np.exp(mean)
    return original_mean, original_mean * std


def build_gp(
    dimensions: int, function_number: int, seed: int, restarts: int
) -> GaussianProcessRegressor:
    signal = ConstantKernel(1.0, constant_value_bounds=(1e-3, 1e3))
    ard_rbf = RBF(
        length_scale=np.full(dimensions, 0.20),
        length_scale_bounds=(0.01, 2.0),
    )
    noise = (
        WhiteKernel(noise_level=0.05, noise_level_bounds=(1e-5, 2.0))
        if function_number == 2
        else WhiteKernel(noise_level=1e-5, noise_level_bounds=(1e-8, 0.10))
    )
    return GaussianProcessRegressor(
        kernel=signal * ard_rbf + noise,
        alpha=1e-10,
        normalize_y=True,
        n_restarts_optimizer=restarts,
        random_state=seed,
    )


def diagnose_trend(y: np.ndarray, function_number: int) -> str:
    if len(y) < 10:
        return "too sparse to establish a reliable trend"
    recent = y[-6:]
    previous = y[-12:-6] if len(y) >= 12 else y[:-6]
    scale = max(float(np.std(y)), 1e-12)
    best_before = float(np.max(y[:-6])) if len(y) > 6 else -np.inf
    if function_number == 2 and float(np.ptp(recent)) > 0.35 * scale:
        return "noisy and recently worsening"
    if float(np.max(recent)) > best_before + 0.05 * scale:
        return "improving, but with variation around the best region"
    if previous.size and float(np.mean(recent)) < float(np.mean(previous)) - 0.20 * scale:
        return "worsening after earlier stronger observations"
    if float(np.ptp(recent)) <= 0.10 * scale:
        return "plateauing near the current region"
    return "mixed, with no stable monotonic trend"


def choose_policy(
    function_number: int,
    trend: str,
    current_best: float,
    external_benchmark: float | None,
) -> Policy:
    """Use supplied trends and benchmarks alongside fixed function rules.

    Reassess rules after each completed round; no budget schedule is applied.
    Round 9 review (2026-09-29): Function 5 improved to 4367.74, above
    its 2883.2 benchmark, supporting the posterior-mean branch below.
    """
    benchmark_gap = external_benchmark is not None and external_benchmark > current_best
    if function_number == 1:
        return Policy("ucb", 3.0, "exploratory", "uncertainty dominates an uninformative surface")
    if function_number == 2:
        return Policy("ei", 0.01, "mixture", "noise-aware near-best validation with EI")
    if function_number == 3 and benchmark_gap:
        return Policy("ucb", 2.5, "exploratory", "a better benchmark suggests another basin")
    if function_number == 5 and benchmark_gap:
        return Policy("ucb", 2.0, "mixture", "a better benchmark warrants broader search")
    if function_number == 5:
        return Policy("mean", 0.0, "exploitative", "no external benchmark exceeds the observed best; posterior-mean refinement is selected")
    if function_number == 4 and "worsening" in trend:
        return Policy("ucb", 1.25, "mixture", "recent deterioration warrants exploration")
    if function_number in (7, 8) and "plateauing" in trend:
        return Policy("mean", 0.0, "exploitative", "the stable local basin supports refinement")
    if function_number in (7, 8):
        return Policy("ucb", 0.75, "mixture", "high dimensionality warrants cautious exploration")
    return Policy("ei", 0.005, "exploitative", "the data support local refinement")


def candidate_pool(
    x: np.ndarray,
    y: np.ndarray,
    function_number: int,
    rng: np.random.Generator,
    seed: int,
    n_global: int,
    n_local: int,
) -> np.ndarray:
    dimensions = x.shape[1]
    exponent = int(np.ceil(np.log2(max(n_global, 2))))
    global_points = qmc.Sobol(d=dimensions, scramble=True, seed=seed).random_base2(exponent)[:n_global]
    batches = [global_points]
    top_indices = np.argsort(y)[-min(3, len(y)):]
    per_centre = max(1, n_local // len(top_indices))
    for radius in LOCAL_RADII:
        for index in top_indices:
            local = x[index] + rng.normal(0.0, radius, size=(per_centre, dimensions))
            batches.append(np.clip(local, 0.0, 1.0))
    if function_number == 2:
        incumbent = x[np.argmax(y)]
        batches.append(np.clip(incumbent + rng.normal(0.0, 0.005, (n_local, dimensions)), 0.0, 1.0))
    candidates = np.unique(np.round(np.vstack(batches), 12), axis=0)
    duplicate = np.zeros(len(candidates), dtype=bool)
    for observed in x:
        duplicate |= np.all(np.isclose(candidates, observed, rtol=0.0, atol=1e-9), axis=1)
    return candidates[~duplicate]


def expected_improvement(
    mean: np.ndarray, std: np.ndarray, incumbent: float, xi: float
) -> np.ndarray:
    improvement = mean - incumbent - xi
    score = np.zeros_like(improvement)
    mask = std > 1e-12
    z = improvement[mask] / std[mask]
    score[mask] = improvement[mask] * norm.cdf(z) + std[mask] * norm.pdf(z)
    return score


def restrict_to_radius(
    candidates: np.ndarray,
    mean: np.ndarray,
    std: np.ndarray,
    incumbent: np.ndarray,
    radius: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    mask = np.linalg.norm(candidates - incumbent, axis=1) <= radius
    return (candidates[mask], mean[mask], std[mask]) if np.any(mask) else (candidates, mean, std)


def recommend(
    x: np.ndarray,
    y: np.ndarray,
    function_number: int,
    rng: np.random.Generator,
    seed: int,
    n_global: int,
    n_local: int,
    restarts: int,
) -> Recommendation:
    transformed_y = transform_output(y, function_number)
    gp = build_gp(x.shape[1], function_number, seed, restarts)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", ConvergenceWarning)
        gp.fit(x, transformed_y)
    trend = diagnose_trend(y, function_number)
    policy = choose_policy(
        function_number,
        trend,
        float(np.max(y)),
        EXTERNAL_OUTPUT_BENCHMARKS[f"function_{function_number}"],
    )
    candidates = candidate_pool(x, y, function_number, rng, seed, n_global, n_local)
    mean, std = gp.predict(candidates, return_std=True)
    incumbent = x[np.argmax(y)]
    if function_number == 5 and policy.acquisition == "mean":
        candidates, mean, std = restrict_to_radius(candidates, mean, std, incumbent, 0.10)
    if function_number in (7, 8) and policy.acquisition == "mean":
        candidates, mean, std = restrict_to_radius(candidates, mean, std, incumbent, 0.25)
    if function_number == 2 and "noisy" in trend:
        candidates, mean, std = restrict_to_radius(candidates, mean, std, incumbent, 0.02)
    if policy.acquisition == "ucb":
        score = mean + policy.exploration * std
    elif policy.acquisition == "mean":
        score = mean
    else:
        score = expected_improvement(mean, std, float(np.max(transformed_y)), policy.exploration)
    chosen = int(np.argmax(score))
    original_mean, original_std = inverse_prediction(mean[[chosen]], std[[chosen]], function_number)
    predicted_mean = float(original_mean[0])
    return Recommendation(
        point=candidates[chosen],
        predicted_mean=predicted_mean,
        predicted_std=float(original_std[0]),
        policy=policy,
        trend=trend,
        kernel=str(gp.kernel_),
        caveat=(
            "surrogate prediction exceeds every observed output; evaluation is required"
            if predicted_mean > float(np.max(y))
            else "none"
        ),
    )


def append_recommendations(
    workbook: Path, recommendations: dict[str, Recommendation]
) -> dict[str, str]:
    from openpyxl import load_workbook

    book = load_workbook(workbook)
    statuses: dict[str, str] = {}
    changed = False
    for sheet, recommendation in recommendations.items():
        worksheet = book[sheet]
        values = [float(value) for value in recommendation.point]
        output_column = len(values) + 1
        if worksheet.cell(worksheet.max_row, output_column).value is None:
            statuses[sheet] = "withheld: prior row has a blank output"
            continue
        duplicate = any(
            all(
                cell.value is not None
                and np.isclose(float(cell.value), values[index], rtol=0.0, atol=1e-9)
                for index, cell in enumerate(row[: len(values)])
            )
            for row in worksheet.iter_rows(min_row=2, max_col=len(values))
        )
        if duplicate:
            statuses[sheet] = "withheld: duplicate input vector"
            continue
        previous_row = worksheet.max_row
        next_row = previous_row + 1
        for column in range(1, output_column + 1):
            source = worksheet.cell(previous_row, column)
            target = worksheet.cell(next_row, column)
            if source.has_style:
                target._style = copy(source._style)
            target.number_format = source.number_format
        for column, value in enumerate(values, start=1):
            worksheet.cell(next_row, column, value)
        worksheet.cell(next_row, output_column, None)
        statuses[sheet] = f"appended to row {next_row}"
        changed = True
    if changed:
        book.save(workbook)
    return statuses


def parse_args() -> argparse.Namespace:
    repository = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description="Generate scikit-learn BO candidates.")
    parser.add_argument("--workbook", type=Path, default=repository / "data" / "Capstone.xlsx")
    append_group = parser.add_mutually_exclusive_group()
    append_group.add_argument("--append", dest="append", action="store_true")
    append_group.add_argument("--no-append", dest="append", action="store_false")
    parser.set_defaults(append=True)
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    parser.add_argument("--n-global", type=int, default=DEFAULT_GLOBAL_CANDIDATES)
    parser.add_argument("--n-local", type=int, default=DEFAULT_LOCAL_CANDIDATES)
    parser.add_argument("--restarts", type=int, default=8)
    args = parser.parse_args()
    if args.n_global < 2 or args.n_local < 1 or args.restarts < 0:
        parser.error("candidate counts must be positive and restarts non-negative")
    return args


def main() -> None:
    args = parse_args()
    workbook = args.workbook.resolve()
    if not workbook.exists():
        raise FileNotFoundError(workbook)
    rng = np.random.default_rng(args.seed)
    recommendations: dict[str, Recommendation] = {}
    reports: list[tuple[str, str]] = []
    for function_number in range(1, 9):
        sheet = f"function_{function_number}"
        data = pd.read_excel(workbook, sheet_name=sheet)
        input_columns = [column for column in data.columns if column.startswith("input_")]
        observed = data.dropna(subset=["output"])
        if observed.empty:
            raise ValueError(f"{sheet} has no completed outputs.")
        x = observed[input_columns].to_numpy(dtype=float)
        y = observed["output"].to_numpy(dtype=float)
        if np.any((x < 0.0) | (x > 1.0)):
            raise ValueError(f"{sheet} contains an input outside [0, 1].")
        result = recommend(
            x, y, function_number, rng, args.seed + function_number,
            args.n_global, args.n_local, args.restarts,
        )
        recommendations[sheet] = result
        best_index = int(np.argmax(y))
        best_input = "-".join(f"{value:.6f}" for value in x[best_index])
        candidate = "-".join(f"{value:.6f}" for value in result.point)
        last_blank = bool(pd.isna(data["output"].iloc[-1]))
        reports.append((sheet,
            f"current_best={y[best_index]:.10g} at {best_input} | next={candidate} | "
            f"predicted={result.predicted_mean:.10g} +/- {result.predicted_std:.4g} | "
            f"strategy={result.policy.label} ({result.policy.acquisition.upper()}) | "
            f"trend={result.trend} | caveat={result.caveat} | "
            f"status={'withheld: prior row unfinished' if last_blank else 'eligible to append'}"
        ))
        print(f"{sheet}: kernel={result.kernel}")
    statuses = append_recommendations(workbook, recommendations) if args.append else {}
    print("\nRecommendations")
    for sheet, report in reports:
        if args.append:
            report += f" | append_result={statuses[sheet]}"
        print(f"{sheet}: {report}")


if __name__ == "__main__":
    main()
