"""Propose the next maximisation input for every black-box function.

This is deliberately self-contained: it uses only numpy and pandas, which makes
it runnable with the bundled workspace Python runtime.  It fits a small ensemble
of RBF Gaussian-process surrogates selected by leave-one-out error, searches the
full [0, 1]^d domain, and prints one input vector per workbook sheet.

Run normally to print and append recommendations. A candidate is appended only
when the preceding row has an output; unfinished rows must be evaluated first.
Use --no-append for a preview-only run that leaves the workbook unchanged.
"""

import argparse
from copy import copy
from pathlib import Path
from datetime import datetime
import numpy as np
import pandas as pd


SEED = int(datetime.now().timestamp())
LENGTH_SCALES = (0.04, 0.06, 0.08, 0.10, 0.14, 0.18, 0.25, 0.35, 0.50)
N_GLOBAL_CANDIDATES = 100_000
N_LOCAL_CANDIDATES = 30_000

# Enter externally reported output values here.  Without corresponding inputs,
# they cannot train the surrogate; they are retained as a best-known benchmark.
OTHER_OBSERVED_OUTPUTS = {
    "function_1": None,
    "function_2": None,
    "function_3": -0.00551,
    "function_4": None,
    "function_5": 2883.2,
    "function_6": None,
    "function_7": None,
    "function_8": 9.939,
}


def fit_rbf_gp(x: np.ndarray, y: np.ndarray, length_scale: float):
    """Fit a noiseless, standardised RBF GP and return its LOO loss."""
    y_mean = float(y.mean())
    y_scale = float(y.std()) or 1.0
    z = (y - y_mean) / y_scale
    squared_distance = ((x[:, None, :] - x[None, :, :]) ** 2).sum(axis=2)
    kernel = np.exp(-0.5 * squared_distance / length_scale**2)
    # A small nugget prevents numerical failure without obscuring the samples.
    inverse_kernel = np.linalg.inv(kernel + np.eye(len(x)) * 1e-6)
    alpha = inverse_kernel @ z
    loo_loss = float(np.mean((alpha / np.diag(inverse_kernel)) ** 2))
    return y_mean, y_scale, alpha, inverse_kernel, loo_loss


def predict_rbf_gp(x: np.ndarray, model, candidates: np.ndarray, length_scale: float):
    """Return posterior mean and standard deviation at candidate points."""
    y_mean, y_scale, alpha, inverse_kernel, _ = model
    squared_distance = ((candidates[:, None, :] - x[None, :, :]) ** 2).sum(axis=2)
    cross_kernel = np.exp(-0.5 * squared_distance / length_scale**2)
    mean = y_mean + y_scale * (cross_kernel @ alpha)
    variance = 1.0 - np.einsum("ij,jk,ik->i", cross_kernel, inverse_kernel, cross_kernel)
    std = y_scale * np.sqrt(np.maximum(variance, 1e-12))
    return mean, std


def candidates_around_best(x: np.ndarray, y: np.ndarray, rng: np.random.Generator) -> np.ndarray:
    """Use global coverage plus increasingly fine local refinement."""
    batches = [rng.random((N_GLOBAL_CANDIDATES, x.shape[1]))]
    centre = x[np.argmax(y)]
    for radius in (0.20, 0.10, 0.05, 0.025, 0.01):
        batches.append(np.clip(centre + rng.normal(0, radius, (N_LOCAL_CANDIDATES, x.shape[1])), 0, 1))
    return np.vstack(batches)


def recommend(x: np.ndarray, y: np.ndarray, function_number: int, rng: np.random.Generator):
    # A log scale prevents the very large values of function 5 from dominating
    # the GP fit, while retaining the location of a maximum.
    transformed_y = np.log(y) if function_number == 5 else y
    fitted = [fit_rbf_gp(x, transformed_y, scale) for scale in LENGTH_SCALES]
    best_indices = np.argsort([model[-1] for model in fitted])[:3]
    weights = np.array([1.0 / fitted[index][-1] for index in best_indices])
    weights /= weights.sum()

    candidates = candidates_around_best(x, y, rng)
    means, stds = [], []
    for index in best_indices:
        mean, std = predict_rbf_gp(x, fitted[index], candidates, LENGTH_SCALES[index])
        means.append(mean)
        stds.append(std)
    mean = np.average(means, axis=0, weights=weights)
    std = np.average(stds, axis=0, weights=weights)

    # Evidence reviewed after round 9 (2026-09-29): Function 1 remains nearly
    # zero; Function 3's best -0.00731434 trails the benchmark -0.00551.
    # Their UCB settings preserve uncertainty-aware exploration.
    # Function 5 improved from 3941.01 to 4367.74, exceeding its 2883.2
    # benchmark. Posterior-mean exploitation is therefore intentional here,
    # as it is for Functions 2, 4, 6, 7 and 8 below.
    # Reassess these fixed settings after each completed round. This function
    # does not read the evaluation budget or automatically adapt its policy.
    if function_number == 1:
        score = mean + 2.0 * std
    elif function_number == 3:
        score = mean + 1.5 * std
    else:
        score = mean
    point = candidates[np.argmax(score)]
    predicted = float(mean[np.argmax(score)])
    if function_number == 5:
        predicted = float(np.exp(predicted))
    return point, predicted


def append_recommendations(workbook: Path, recommendations: dict[str, np.ndarray]) -> None:
    """Append unseen input vectors, retaining the preceding row's formatting."""
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("--append requires openpyxl. Install it in the Python environment first.") from error

    book = load_workbook(workbook)
    for sheet, point in recommendations.items():
        worksheet = book[sheet]
        values = [float(value) for value in point]
        input_count = len(values)
        if worksheet.cell(worksheet.max_row, input_count + 1).value is None:
            print(f"{sheet}: last output is blank; not appended")
            continue
        already_present = any(
            all(
                cell.value is not None and np.isclose(float(cell.value), values[index])
                for index, cell in enumerate(row[:input_count])
            )
            for row in worksheet.iter_rows(min_row=2, max_col=input_count)
        )
        if already_present:
            print(f"{sheet}: input vector already present; not appended")
            continue

        previous_row = worksheet.max_row
        next_row = previous_row + 1
        # Copy the prior row's style so the new evaluation row matches the sheet.
        for column in range(1, input_count + 2):
            source = worksheet.cell(previous_row, column)
            target = worksheet.cell(next_row, column)
            if source.has_style:
                target._style = copy(source._style)
            if source.number_format:
                target.number_format = source.number_format
        for column, value in enumerate(values, start=1):
            worksheet.cell(next_row, column, value)
        worksheet.cell(next_row, input_count + 1, None)
        print(f"{sheet}: appended inputs to row {next_row}")
    book.save(workbook)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate Bayesian-optimisation candidates for Capstone.xlsx.")
    append_group = parser.add_mutually_exclusive_group()
    append_group.add_argument(
        "--append",
        dest="append",
        action="store_true",
        help="append each generated input vector (the default)",
    )
    append_group.add_argument(
        "--no-append",
        dest="append",
        action="store_false",
        help="print recommendations without changing the workbook",
    )
    parser.set_defaults(append=True)
    parser.add_argument(
        "--seed",
        type=int,
        default=SEED,
        help="random seed for reproducible candidate generation",
    )
    args = parser.parse_args()
    workbook = Path(__file__).with_name("Capstone.xlsx")
    rng = np.random.default_rng(args.seed)
    recommendations: dict[str, np.ndarray] = {}
    for function_number in range(1, 9):
        sheet = f"function_{function_number}"
        data = pd.read_excel(workbook, sheet_name=sheet)
        input_columns = [column for column in data.columns if column.startswith("input_")]
        observed = data.dropna(subset=["output"])
        x = observed[input_columns].to_numpy(dtype=float)
        y = observed["output"].to_numpy(dtype=float)
        point, predicted = recommend(x, y, function_number, rng)
        recommendations[sheet] = point
        current_best = float(y.max())
        reported_best = OTHER_OBSERVED_OUTPUTS.get(sheet)
        best_known = current_best if reported_best is None else max(current_best, float(reported_best))
        print(
            f"{sheet}: "
            f"inputs=[{'-'.join(f'{value:.6f}' for value in point)}]  "
            f"surrogate_prediction={predicted:.6g}  "
            f"current_best={current_best:.6g}  "
            f"best_known={best_known:.6g}  "
            f"last_output={'blank' if pd.isna(data['output'].iloc[-1]) else 'present'}"
        )
    if args.append:
        append_recommendations(workbook, recommendations)


if __name__ == "__main__":
    main()
