"""Run with: python -m unittest discover -s tests"""
import json
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from capstone_rounds import append_round, refresh_repository
from openpyxl import Workbook, load_workbook


class AppendBudgetTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.root = Path(self.folder.name)
        self.workbook = self.root / "Capstone.xlsx"
        self.ledger = self.root / "evaluation_budget.json"
        self.names = [f"function_{i}" for i in range(1, 9)]
        book = Workbook()
        book.remove(book.active)
        for name in self.names:
            sheet = book.create_sheet(name)
            sheet.append(["input_1", "input_2", "output"])
            sheet.append([0.1, 0.2, 1.0])
        book.save(self.workbook)
        book.close()
        self.budget = {
            "used_rounds": 1, "remaining_rounds": 12, "total_rounds": 13,
            "round_log": [{"round": 1, "status": "completed",
                           "workbook_rows": dict.fromkeys(self.names, 2)}],
        }
        self.ledger.write_text(json.dumps(self.budget), encoding="utf-8")
        self.points = {n: SimpleNamespace(point=[0.3, 0.4]) for n in self.names}

    def snapshot(self):
        return self.workbook.read_bytes(), self.ledger.read_bytes()

    def test_append_charges_once_and_refreshes_repository(self):
        (self.root / "repository/data").mkdir(parents=True)
        result = append_round(self.workbook, self.points)
        self.assertTrue(all("round 2 charged" in value for value in result.values()))
        ledger = json.loads(self.ledger.read_text())
        self.assertEqual((ledger["used_rounds"], ledger["remaining_rounds"]), (2, 11))
        self.assertEqual(ledger["round_log"][-1]["workbook_rows"], dict.fromkeys(self.names, 3))
        book = load_workbook(self.workbook)
        for sheet in book:
            self.assertEqual(list(sheet.values)[1], (0.1, 0.2, 1))
            self.assertEqual(list(sheet.values)[2], (0.3, 0.4, None))
        book.close()
        before = self.snapshot()
        append_round(self.workbook, self.points)
        self.assertEqual(before, self.snapshot())
        refresh_repository(self.workbook)
        self.assertEqual(before, self.snapshot())
        for path in [self.workbook, self.ledger]:
            self.assertEqual(path.read_bytes(), (self.root / "repository/data" / path.name).read_bytes())

    def test_duplicate_blocks_entire_round(self):
        self.points[self.names[-1]].point = [0.1, 0.2]
        before = self.snapshot()
        result = append_round(self.workbook, self.points)
        self.assertTrue(all("withheld" in value for value in result.values()))
        self.assertEqual(before, self.snapshot())

    def test_single_pending_sheet_blocks_entire_round(self):
        book = load_workbook(self.workbook)
        book[self.names[-1]]["C2"] = None
        book.save(self.workbook)
        book.close()
        before = self.snapshot()
        append_round(self.workbook, self.points)
        self.assertEqual(before, self.snapshot())

    def test_exhausted_budget_and_partial_round_do_not_write(self):
        self.budget.update(total_rounds=1, remaining_rounds=0)
        self.ledger.write_text(json.dumps(self.budget), encoding="utf-8")
        before = self.snapshot()
        result = append_round(self.workbook, self.points)
        self.assertTrue(all("exhausted" in value for value in result.values()))
        self.assertEqual(before, self.snapshot())
        self.points.pop(self.names[-1])
        with self.assertRaises(ValueError):
            append_round(self.workbook, self.points)
        self.assertEqual(before, self.snapshot())

    def test_ledger_write_failure_restores_workbook(self):
        before = self.snapshot()
        replace = Path.replace
        def fail_ledger(path, target):
            if Path(target) == self.ledger:
                raise PermissionError("simulated ledger lock")
            return replace(path, target)
        with patch.object(Path, "replace", fail_ledger):
            with self.assertRaises(PermissionError):
                append_round(self.workbook, self.points)
        self.assertEqual(before, self.snapshot())


if __name__ == "__main__":
    unittest.main()
